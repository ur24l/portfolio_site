"""Build the resume from the same verified profile used by the website.
The layout uses ordinary paragraphs and tab stops, not tables or text boxes.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
p=json.loads((ROOT/'src/data/profile.json').read_text())
doc=Document()
section=doc.sections[0]
section.page_width,section.page_height=Inches(8.27),Inches(11.69)
section.top_margin=Inches(.48)
section.bottom_margin=Inches(.48)
section.left_margin=section.right_margin=Inches(.6)
text_width=section.page_width-section.left_margin-section.right_margin
normal=doc.styles['Normal']
normal.font.name='Arial'
normal.font.size=Pt(10.5)
normal.font.color.rgb=RGBColor.from_string('253044')
normal.paragraph_format.line_spacing=1.06
normal.paragraph_format.space_after=Pt(2)
for name in ['Title','Heading 1','Heading 2']:
 s=doc.styles[name]
 s.font.name='Arial'
 s.font.color.rgb=RGBColor.from_string('172A47')
 s.paragraph_format.keep_with_next=True
 s.paragraph_format.space_after=Pt(4)
 s.paragraph_format.space_before=Pt(10)
 for border in s.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].font.bold=True
doc.styles['Title'].font.color.rgb=RGBColor.from_string('111B2D')
doc.styles['Title'].paragraph_format.space_before=Pt(0)
doc.styles['Title'].paragraph_format.space_after=Pt(3)
doc.styles['Heading 1'].font.size=Pt(10.5)
doc.styles['Heading 1'].font.bold=True

def paragraph(text='',size=None,bold=False,after=2,color=None):
 item=doc.add_paragraph()
 item.paragraph_format.space_after=Pt(after)
 r=item.add_run(text);r.bold=bold
 if size:r.font.size=Pt(size)
 if color:r.font.color.rgb=RGBColor.from_string(color)
 return item

def hyperlink(item,text,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),item.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Arial');font.set(qn('w:hAnsi'),'Arial');pr.append(font)
 size=OxmlElement('w:sz');size.set(qn('w:val'),'19');pr.append(size)
 color=OxmlElement('w:color');color.set(qn('w:val'),'36527C');pr.append(color)
 r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);item._p.append(h)

def heading(text):doc.add_paragraph(text.upper(),'Heading 1')
def role(title,dates):
 item=paragraph(after=1)
 item.paragraph_format.keep_with_next=True
 item.paragraph_format.tab_stops.add_tab_stop(text_width,WD_TAB_ALIGNMENT.RIGHT)
 r=item.add_run(title);r.bold=True
 r=item.add_run('\t'+dates);r.font.size=Pt(9.5);r.font.color.rgb=RGBColor.from_string('596477')
 return item

def bullet(text):
 item=doc.add_paragraph(style='List Bullet')
 item.add_run(text)
 item.paragraph_format.left_indent=Inches(.14)
 item.paragraph_format.first_line_indent=Inches(-.11)
 item.paragraph_format.space_after=Pt(2)
 item.paragraph_format.line_spacing=1.05
 item.paragraph_format.keep_together=True
 return item

doc.add_paragraph(p['name'],'Title')
paragraph('Junior Software Developer',size=12,bold=False,after=5,color='45566E')
contact=paragraph(size=9.5,after=1)
contact.add_run(p['location']+'  |  '+p['phone']+'  |  ').font.size=Pt(9.5)
hyperlink(contact,p['email'],'mailto:'+p['email'])
links=paragraph(after=0)
hyperlink(links,'linkedin.com/in/bhargav-sai-bhartu',p['linkedin'])
links.add_run('  |  ').font.size=Pt(9.5)
hyperlink(links,'github.com/ur24l',p['github'])
heading('Professional Summary')
paragraph(p['summary'])
heading('Experience')
for job in p['experience']:
 role(job['title'],job['dates'].replace('January','Jan').replace('March','Mar').replace('July','Jul'))
 org=paragraph(job['org'],size=10,after=3,color='596477')
 org.paragraph_format.keep_with_next=True
 for line in job['bullets']:bullet(line)
heading('Skills')
for label,items in p['skills']:
 item=paragraph(after=2)
 item.add_run(label+': ').bold=True
 item.add_run(items)
heading('Selected Projects')
for project in p['projects']:
 item=paragraph(after=1)
 item.paragraph_format.keep_with_next=True
 item.add_run(project['name']).bold=True
 r=item.add_run('  |  '+project['tech']);r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('596477')
 for line in project['bullets']:bullet(line)
heading('Education')
for e in p['education']:
 degree=e['degree'].replace('Bachelor of Technology (B.Tech),','B.Tech in').replace('Diploma,','Diploma in')
 role(degree,e['dates'])
 paragraph(e['school']+' | '+e['detail'],size=9.5,after=3,color='596477')
heading('Achievements and Certifications')
for line in p['achievements']:paragraph(line,size=9.5,after=2)
paragraph('Languages: Telugu (native), English (fluent)',size=9.5,after=0,color='596477')
doc.core_properties.title=p['name']+' Resume'
doc.core_properties.author=p['name']
doc.core_properties.subject='Junior Software Developer'
out=ROOT/'public/assets/Bhargav-Sai-Bhartu-Resume.docx'
doc.save(out)
print(out)
